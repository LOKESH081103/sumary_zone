#!/usr/bin/env python3
"""
petition.py   (Project 2 - Section 138 COMPLAINT PETITIONS)
============================================================
1. Reads every complaint-petition PDF in a folder.
2. Extracts the accused parties (name + address, incl. "ALSO AT" addresses),
   the cheque table (date / no / bank / amount / total), the return reason and
   the notice / filing dates.
3. The petition carries NO agreement number, so the matching company.xlsx row is
   found from: cheque no + cheque amount + cheque date + accused names.
   (If an agreement number does appear in the PDF, that is used directly.)
4. Compares everything with company.xlsx and writes
   petition_reconciliation_report.xlsx with an OK / CHECK verdict per petition.

USAGE
-----
    python petition.py                                   # PDFs + company.xlsx in this folder
    python petition.py --pdf-folder D:\\petitions --company D:\\company.xlsx

INSTALL
-------
    pip install pandas openpyxl pdfplumber rapidfuzz     (rapidfuzz optional)
"""

import argparse
import os
import re
import sys
import time
from datetime import date, datetime

import pandas as pd

try:
    import pdfplumber
except ImportError:  # pragma: no cover
    pdfplumber = None

# ----------------------------------------------------------------------------
# CONFIG
# ----------------------------------------------------------------------------
MATCH_T = 90        # name AND address >= this   -> MATCH
PARTIAL_T = 60      # name AND address >= this   -> PARTIAL MATCH, else MISMATCH
PAIR_MIN_NAME = 50  # accused is only paired with a company party if names are this similar

# company-row lookup (petitions have no agreement number)
W_CHEQUE_NO = 40
W_AMOUNT = 25
W_DATE = 10
W_NAMES = 25
ACCEPT_MIN_SCORE = 30   # below this -> NOT FOUND
HIGH_MIN_SCORE = 60     # HIGH confidence needs this score ...
HIGH_MIN_MARGIN = 15    # ... and this lead over the runner-up

# ----------------------------------------------------------------------------
# Fuzzy similarity (rapidfuzz if present, difflib otherwise)
# ----------------------------------------------------------------------------
try:
    from rapidfuzz import fuzz

    def _ratio(a, b):
        return fuzz.ratio(a, b)

    def _tsort(a, b):
        return fuzz.token_sort_ratio(a, b)

except ImportError:
    from difflib import SequenceMatcher

    def _ratio(a, b):
        return 100.0 * SequenceMatcher(None, a, b).ratio()

    def _tsort(a, b):
        return _ratio(" ".join(sorted(a.split())), " ".join(sorted(b.split())))


def isnull(v):
    if v is None:
        return True
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def norm(text):
    """Upper-case, punctuation -> space, collapse whitespace."""
    if isnull(text):
        return ""
    return re.sub(r"[^A-Z0-9]+", " ", str(text).upper()).strip()


def similarity(a, b):
    """0-100. Ignores spaces (company data often glues words together) and word order."""
    a, b = norm(a), norm(b)
    if not a and not b:
        return 100.0
    if not a or not b:
        return 0.0
    return max(_ratio(a.replace(" ", ""), b.replace(" ", "")), _tsort(a, b))


def parse_date(v):
    if isnull(v) or str(v).strip() == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = re.sub(r"\s+", "", str(v))
    m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})$", s)
    if m:
        d, mo, y = (int(x) for x in m.groups())
        if y < 100:
            y += 2000
        try:
            return date(y, mo, d)
        except ValueError:
            return None
    try:
        return pd.to_datetime(s, dayfirst=True).date()
    except Exception:
        return None


def to_number(v):
    if isnull(v):
        return None
    s = re.sub(r"[^\d.]", "", str(v))
    try:
        return float(s) if s else None
    except ValueError:
        return None


def cheque_key(v):
    """'000504' == 504 == 504.0 == '504'"""
    if isnull(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"[^A-Z0-9]", "", str(v).upper()).lstrip("0")


# ----------------------------------------------------------------------------
# PART 1 - PETITION PDF EXTRACTION
# ----------------------------------------------------------------------------
DATE_RX = r"\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}"
# dates in running text can be broken over two lines ("15-09-\n26")
DATE_FLEX = r"\d{1,2}\s*[-/.]\s*\d{1,2}\s*[-/.]\s*\d{2,4}"

ALSO_RE = re.compile(
    r"\b(?:also\s+(?:at|add(?:ress)?|residing\s+at|resides\s+at|r/o)"
    r"|alternate\s+address|another\s+address)\b\s*[:\-\u2013\u2014.]*\s*",
    re.IGNORECASE,
)
NUM_RE = re.compile(r"^(\d+)\s*[.)]\s*(.*)$")
NOISE_RES = [
    re.compile(r"^[A-Z]\s*[A-Z]\s*(?:\d\s*){6,}[A-Z]\s*[A-Z]$", re.I),
    re.compile(r"^DATED\s*:", re.I),
    re.compile(r"^page\s+\d+(\s+of\s+\d+)?$", re.I),
]
AGREEMENT_RE = re.compile(r"\b([A-Z]{2}\d{2}[A-Z]{3}\d{8,})\b")

VERSUS_RE = re.compile(r"^[ \t\-\u2013\u2014]*Versus[ \t\-\u2013\u2014]*$", re.I | re.M)
COMPLAINANT_RE = re.compile(r"Complainant[ \t]*$", re.I | re.M)
ACCUSED_RE = re.compile(r"^[ \t.\u2026\u00b7_\-]*Accused\b", re.I | re.M)
PETITION_START_RE = re.compile(r"^[ \t]*(?:IN\s+THE\s+)+COURT\s+OF\b", re.I | re.M)


def is_noise(line):
    return any(rx.search(line) for rx in NOISE_RES)


def read_pdf_text(path):
    if pdfplumber is None:
        sys.exit("pdfplumber is missing:  pip install pdfplumber")
    chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                chunks.append(t)
    return "\n".join(chunks).replace("\r\n", "\n").replace("\r", "\n")


def looks_like_petition(text):
    return bool(re.search(r"\bVersus\b", text, re.I) and re.search(r"\bAccused\b", text, re.I))


def split_petitions(text):
    """A PDF normally holds one petition; if several are stacked, split them."""
    starts = [m.start() for m in PETITION_START_RE.finditer(text)]
    if len(starts) <= 1:
        return [text]
    starts.append(len(text))
    return [text[a:b] for a, b in zip(starts[:-1], starts[1:])]


def parse_parties(block):
    """
    [{"num", "name", "addresses": [...]}]  - numbers must run 1,2,3.. in order, and
    'ALSO AT ...' (own line or inline) starts an extra address for the same party.
    """
    parties, cur, expected = [], None, 1
    for raw in block.split("\n"):
        line = raw.strip()
        if not line or is_noise(line):
            continue
        m = NUM_RE.match(line)
        if m and int(m.group(1)) == expected:
            cur = {"num": expected, "name": m.group(2).strip(), "addr_lines": [[]]}
            parties.append(cur)
            expected += 1
            continue
        if cur is None:
            continue
        if not cur["name"]:
            cur["name"] = line
            continue
        if re.match(r"^\d+\s*[.)]?$", line):
            continue
        first, *rest = ALSO_RE.split(line)
        if first.strip():
            cur["addr_lines"][-1].append(first.strip())
        for part in rest:
            cur["addr_lines"].append([part.strip()] if part.strip() else [])

    out = []
    for p in parties:
        addrs = [", ".join(a) for a in p["addr_lines"]]
        addrs = [a for i, a in enumerate(addrs) if a or i == 0]
        if p["name"] and len(p["name"]) >= 2:
            out.append({"num": p["num"], "name": p["name"], "addresses": addrs})
    return out


def accused_block(text):
    """Text between '- Versus -' and '.... Accused'."""
    acc = ACCUSED_RE.search(text)
    if not acc:
        return ""
    head = text[: acc.start()]
    vs = None
    for vs in VERSUS_RE.finditer(head):
        pass                                   # keep the last 'Versus' before 'Accused'
    if vs:
        return head[vs.end():]
    comp = None
    for comp in COMPLAINANT_RE.finditer(head):
        pass
    return head[comp.end():] if comp else ""


def parse_cheque_table(text):
    """
    CHEQUE DATE  CHEQUE NO  DRAWN ON  AMOUNT
    07-09-26     910330     INDUSIND BANK  2497808/-
    TOTAL 2497808/-
    Returns (list of cheques, total string)
    """
    cheques, total = [], ""
    m = re.search(
        r"CHEQUE\s+DATE\s+CHEQUE\s+NO\.?\s+DRAWN\s+ON\s+AMOUNT(.*?)(?=^[ \t]*TOTAL\b|\Z)",
        text, re.I | re.S | re.M,
    )
    if m:
        row_rx = re.compile(
            r"^[ \t]*(" + DATE_RX + r")[ \t]+([A-Z0-9\-/]+)[ \t]+(.+?)[ \t]+"
            r"(?:Rs\.?[ \t]*)?([\d,]+(?:\.\d+)?)[ \t]*(?:/-|/\-)?[ \t]*$",
            re.I | re.M,
        )
        for r in row_rx.finditer(m.group(1)):
            cheques.append({
                "date": r.group(1), "no": r.group(2).strip(),
                "bank": r.group(3).strip(), "amount": r.group(4).replace(",", ""),
            })
    t = re.search(r"^[ \t]*TOTAL\s*[:\-]?\s*(?:Rs\.?)?\s*([\d,]+(?:\.\d+)?)", text, re.I | re.M)
    if t:
        total = t.group(1).replace(",", "")
    return cheques, total


def parse_cheque_paragraph(text):
    """Fallback for petitions that describe the cheque in a sentence instead of a table."""
    m = re.search(
        r"cheque\s+(?:bearing\s+)?no\.?\s*[:\-]?\s*([\w\-/]+)\s*,?\s*dated\s*[-:\s]*(" + DATE_RX + r")"
        r"\s+for\s+a\s+sum\s+of\s+(?:Rs\.?|INR|\u20b9)\s*([\d,]+(?:\.\d+)?)(?:/-)?"
        r"\s+drawn\s+on\s+(.*?)(?=\s+in\s+favou?r|\s+vide|\.(?:\s|$)|\n|$)",
        text, re.I | re.S,
    )
    if not m:
        return [], ""
    no, dt, amt, bank = (g.strip() for g in m.groups())
    amt = amt.replace(",", "")
    return [{"date": dt, "no": no, "bank": bank, "amount": amt}], amt


def parse_petition_text(text, filename="", index=1):
    """Pure-text parser (unit-testable without a PDF). Returns one petition dict."""
    flat = re.sub(r"\s+", " ", text)

    cheques, total = parse_cheque_table(text)
    if not cheques:
        cheques, total = parse_cheque_paragraph(text)

    m = re.search(r"dishonou?red[^\u201c\u201d\"]{0,120}?[\u201c\"]\s*([^\u201d\"]+?)\s*[\u201d\"]", flat, re.I)
    if m:
        reason = m.group(1).strip().rstrip(".")
    else:
        m = re.search(r"\b([A-Z]\d{3}\s*[-\u2013]\s*[A-Z][A-Z /]+?)(?=\s+(?:vide|in\b|and\b)|[.,]|$)", flat)
        reason = m.group(1).strip() if m else ""

    def grab(rx):
        mm = re.search(rx, flat, re.I)
        return re.sub(r"\s+", "", mm.group(1)) if mm else ""

    memo_date = grab(r"memo(?:/s)?\s+dated\s*(" + DATE_FLEX + ")")
    notice_date = grab(r"demand\s+notice\s+through\s+its\s+.{0,60}?\bdated\s*(" + DATE_FLEX + ")")
    posted_on = grab(r"posted\s+on\s*(" + DATE_FLEX + ")")
    m = re.search(r"This\s+(\d{1,2})\w{0,2}\s+day\s+of\s+([A-Za-z]+)\s+(\d{4})", flat, re.I)
    filing_date = f"{m.group(1)} {m.group(2)} {m.group(3)}" if m else ""
    m = re.search(r"SL\s*NO\s*[:.\-]?\s*(\d+)", flat, re.I)
    sl_no = m.group(1) if m else str(index)

    agreement = ""
    for src in (text, filename):
        mm = AGREEMENT_RE.search(src or "")
        if mm:
            agreement = mm.group(1).upper()
            break

    return {
        "source": os.path.basename(filename), "sl_no": sl_no, "agreement": agreement,
        "parties": parse_parties(accused_block(text)),
        "cheques": cheques, "total": total, "reason": reason,
        "memo_date": memo_date, "notice_date": notice_date,
        "posted_on": posted_on, "filing_date": filing_date,
    }


def extract_petitions_from_pdf(path):
    text = read_pdf_text(path)
    if not looks_like_petition(text):
        return None                                    # not a petition (e.g. a demand notice)
    segs = split_petitions(text)
    return [parse_petition_text(s, path, i) for i, s in enumerate(segs, start=1)]


def flatten(petitions):
    """One row per party-address, for the 'Extracted' sheet."""
    rows = []
    for p in petitions:
        ch = p["cheques"]
        base = {
            "Source File": p["source"], "SL No": p["sl_no"], "Agreement No (in PDF)": p["agreement"],
            "Cheque No": "; ".join(c["no"] for c in ch),
            "Cheque Date": "; ".join(c["date"] for c in ch),
            "Cheque Amount": "; ".join(c["amount"] for c in ch),
            "Drawn On": "; ".join(c["bank"] for c in ch),
            "Cheques Total": p["total"], "Return Reason": p["reason"],
            "Memo Date": p["memo_date"], "Notice Date": p["notice_date"],
            "Posted On": p["posted_on"], "Filing Date": p["filing_date"],
        }
        for party in p["parties"]:
            for seq, addr in enumerate(party["addresses"], start=1):
                rows.append({**{"Source File": base["Source File"], "SL No": base["SL No"],
                                "Party Number": party["num"], "Address Seq": seq,
                                "Name": party["name"], "Address": addr}, **base})
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------
# PART 2 - COMPANY DATA (wide -> list of parties per agreement)
# ----------------------------------------------------------------------------
def load_company(path):
    df = pd.read_excel(path)
    cols = {re.sub(r"[^A-Z0-9]", "", str(c).upper()): c for c in df.columns}

    def pick(*names):
        for n in names:
            if n in cols:
                return cols[n]
        return None

    c_agr = pick("AGREEMENTNO", "AGREEMENTNUMBER", "LOANACCOUNTNO", "LOANACCOUNTNUMBER")
    c_chq = pick("CHEQUENO", "CHEQUENUMBER")
    c_cdt = pick("CHEQUEDATE")
    c_amt = pick("CLAIMAMOUNT", "CHEQUEAMOUNT")
    c_bank = pick("BANK", "DRAWNON")
    c_rsn = pick("CHEQUEREALISTAION", "CHEQUEREALISATION", "RETURNREASON")
    c_bn = pick("BORROWER", "BORROWERNAME")
    c_ba = pick("BORROWERADD", "BORROWERADDRESS")
    if not (c_agr and c_bn and c_ba):
        raise ValueError(f"Agreement / Borrower columns not found. Columns: {list(df.columns)}")

    co = {}
    for sq, orig in cols.items():
        m = re.fullmatch(r"COBORROWER(\d+)(NAME)?", sq)
        if m:
            co.setdefault(int(m.group(1)), {})["name"] = orig
        m = re.fullmatch(r"COBORROWER(\d+)ADD(?:RESS)?", sq)
        if m:
            co.setdefault(int(m.group(1)), {})["add"] = orig
    co = {n: v for n, v in sorted(co.items()) if "name" in v}

    agreements = {}
    for _, r in df.iterrows():
        agr = "" if isnull(r[c_agr]) else str(r[c_agr]).strip().upper()
        if not agr:
            continue
        if agr in agreements:
            print(f"  WARN: agreement {agr} appears more than once in company data - using first row")
            continue
        parties = []
        if not isnull(r[c_bn]) and str(r[c_bn]).strip():
            parties.append({"slot": "Borrower", "name": r[c_bn], "addr": r[c_ba]})
        for n, v in co.items():
            if isnull(r[v["name"]]) or not str(r[v["name"]]).strip():
                continue
            addr = r[v["add"]] if "add" in v else ""
            parties.append({"slot": f"Co-Borrower-{n}", "name": r[v["name"]], "addr": addr})
        g = lambda c: (r[c] if c else None)
        agreements[agr] = {
            "parties": parties, "cheque_no": g(c_chq), "cheque_date": g(c_cdt),
            "amount": g(c_amt), "bank": g(c_bank), "reason": g(c_rsn),
        }
    return agreements


# ----------------------------------------------------------------------------
# PART 3 - FIND THE COMPANY ROW (petition has no agreement number)
# ----------------------------------------------------------------------------
def _name_fraction(pet, comp):
    names = []
    for p in pet["parties"]:
        if p["name"] not in names:
            names.append(p["name"])
    if not names or not comp["parties"]:
        return 0.0
    hit = sum(1 for n in names if max(similarity(n, c["name"]) for c in comp["parties"]) >= 85)
    return hit / len(names)


def score_company(pet, comp, with_names=True):
    score, why = 0.0, []
    ckeys = {cheque_key(c["no"]) for c in pet["cheques"]}
    if cheque_key(comp["cheque_no"]) and cheque_key(comp["cheque_no"]) in ckeys:
        score += W_CHEQUE_NO
        why.append("cheque no")
    amts = [to_number(c["amount"]) for c in pet["cheques"]] + [to_number(pet["total"])]
    ca = to_number(comp["amount"])
    if ca is not None and any(a is not None and abs(a - ca) < 0.5 for a in amts):
        score += W_AMOUNT
        why.append("amount")
    cd = parse_date(comp["cheque_date"])
    if cd and any(parse_date(c["date"]) == cd for c in pet["cheques"]):
        score += W_DATE
        why.append("cheque date")
    if with_names:
        frac = _name_fraction(pet, comp)
        if frac > 0:
            score += W_NAMES * frac
            why.append(f"names {frac:.0%}")
    return score, why


def find_company(pet, companies):
    """-> (agreement, score, matched_by, confidence, note)"""
    if pet["agreement"] and pet["agreement"] in companies:
        return pet["agreement"], 100.0, "Agreement no. in PDF", "HIGH", ""

    strong = [a for a, c in companies.items() if score_company(pet, c, with_names=False)[0] >=
              min(W_CHEQUE_NO, W_AMOUNT)]
    pool = strong if strong else list(companies)       # no cheque/amount hit -> try names across all
    scored = []
    for a in pool:
        s, why = score_company(pet, companies[a])
        scored.append((s, a, why))
    scored.sort(key=lambda x: x[0], reverse=True)
    if not scored:
        return None, 0.0, "", "NONE", "company file is empty"

    best_s, best_a, why = scored[0]
    if best_s < ACCEPT_MIN_SCORE:
        return None, best_s, "", "NONE", f"closest was {best_a} (score {best_s:.0f})"
    margin = best_s - (scored[1][0] if len(scored) > 1 else 0)
    if best_s >= HIGH_MIN_SCORE and margin >= HIGH_MIN_MARGIN:
        conf, note = "HIGH", ""
    else:
        conf = "LOW"
        note = f"score {best_s:.0f}, runner-up {scored[1][1]} ({scored[1][0]:.0f})" if len(scored) > 1 \
            else f"score {best_s:.0f}"
    return best_a, best_s, " + ".join(why), conf, note


# ----------------------------------------------------------------------------
# PART 4 - RECONCILIATION
# ----------------------------------------------------------------------------
def party_status(name_sim, addr_sim):
    if name_sim >= MATCH_T and addr_sim >= MATCH_T:
        return "MATCH"
    if name_sim >= PARTIAL_T and addr_sim >= PARTIAL_T:
        return "PARTIAL MATCH"
    return "MISMATCH"


def match_parties(rows, company_parties):
    """Order-independent: best (name 60% + address 40%) pairs are locked in first."""
    pairs = []
    for i, n in enumerate(rows):
        for j, c in enumerate(company_parties):
            ns = similarity(n["Name"], c["name"])
            a_s = similarity(n["Address"], c["addr"])
            if ns >= PAIR_MIN_NAME:
                pairs.append((0.6 * ns + 0.4 * a_s, i, j, ns, a_s))
    pairs.sort(reverse=True)
    used_i, used_j, assign = set(), set(), {}
    for _, i, j, ns, a_s in pairs:
        if i in used_i or j in used_j:
            continue
        used_i.add(i)
        used_j.add(j)
        assign[i] = (j, ns, a_s)
    return assign


def cheque_checks(pet, comp):
    """-> list of (field, petition_value, company_value, result)"""
    out = []
    ch = pet["cheques"]
    if not ch:
        for f in ("Cheque No", "Cheque Date", "Cheque Amount", "Drawn On / Bank"):
            out.append((f, "", "", "NOT EXTRACTED"))
    else:
        primary = next((c for c in ch if cheque_key(c["no"]) == cheque_key(comp["cheque_no"])), ch[0])
        multi = len(ch) > 1

        nos = ", ".join(c["no"] for c in ch)
        res = "OK" if any(cheque_key(c["no"]) == cheque_key(comp["cheque_no"]) for c in ch) else "MISMATCH"
        out.append(("Cheque No", nos, "" if isnull(comp["cheque_no"]) else comp["cheque_no"], res))

        pd_, cd_ = parse_date(primary["date"]), parse_date(comp["cheque_date"])
        res = "NOT EXTRACTED" if pd_ is None else ("OK" if pd_ == cd_ else "MISMATCH")
        out.append(("Cheque Date", pd_, cd_, res))

        pa, ca = to_number(primary["amount"]), to_number(comp["amount"])
        tot = to_number(pet["total"])
        if pa is None:
            res = "NOT EXTRACTED"
        else:
            ok = ca is not None and (abs(pa - ca) < 0.5 or (multi and tot is not None and abs(tot - ca) < 0.5))
            res = "OK" if ok else "MISMATCH"
        out.append(("Cheque Amount", pa if not multi else f"{pa} (total {tot})", ca, res))

        nb, cb = norm(primary["bank"]), norm(comp["bank"])
        if not nb:
            res = "NOT EXTRACTED"
        else:
            res = "OK" if ((nb in cb or cb in nb) and cb) or similarity(nb, cb) >= 85 else "MISMATCH"
        out.append(("Drawn On / Bank", primary["bank"], "" if isnull(comp["bank"]) else comp["bank"], res))

    nr, cr = pet["reason"], comp["reason"]
    if not nr:
        res = "NOT EXTRACTED"
    else:
        code = lambda s: (re.match(r"\s*([A-Z]\d{3})\b", str(s).upper()) or [None, None])[1]
        nc, cc = code(nr), code(cr)
        res = ("OK" if nc == cc else "MISMATCH") if (nc and cc) else ("OK" if similarity(nr, cr) >= 85 else "MISMATCH")
    out.append(("Return Reason", nr, "" if isnull(cr) else cr, res))
    return out


def _party_row(p, r, agr, c, status, ns, a_s, remark):
    return {
        "Source File": p["source"], "SL No": p["sl_no"], "Agreement No": agr or "",
        "Petition Party No": r["Party Number"], "Address Seq": r["Address Seq"],
        "Status": status, "Petition Name": r["Name"],
        "Company Name": c["name"] if c else "",
        "Name Similarity %": None if ns is None else round(ns, 1),
        "Petition Address": r["Address"], "Company Address": c["addr"] if c else "",
        "Address Similarity %": None if a_s is None else round(a_s, 1),
        "Company Slot": c["slot"] if c else "", "Remarks": remark,
    }


def _remark(status, ns, a_s):
    if status == "MATCH":
        return "Name and address match."
    bits = []
    if ns < MATCH_T:
        bits.append(f"name differs ({ns:.0f}%)")
    if a_s < MATCH_T:
        bits.append(f"address differs ({a_s:.0f}%)")
    return "Check: " + " and ".join(bits) + "."


def reconcile(petitions, companies):
    party_rows, summary_rows = [], []

    for p in petitions:
        rows = []
        for party in p["parties"]:
            for seq, addr in enumerate(party["addresses"], start=1):
                rows.append({"Party Number": party["num"], "Address Seq": seq,
                             "Name": party["name"], "Address": addr})

        agr, score, by, conf, note = find_company(p, companies)
        head = {"Source File": p["source"], "SL No": p["sl_no"]}

        if agr is None:
            for r in rows:
                party_rows.append(_party_row(p, r, "", None, "NOT IN COMPANY DATA", None, None,
                                             "No company row matches this petition's cheque / amount / names."))
            summary_rows.append({**head, "Agreement No": "", "Matched By": "", "Match Confidence": "NONE",
                                 "Verdict": "NOT FOUND",
                                 "Issues": "Could not identify the company row" + (f" ({note})" if note else "")})
            continue

        comp = companies[agr]
        cparties = comp["parties"]
        assign = match_parties(rows, cparties)
        matched_j = {j for j, _, _ in assign.values()}
        issues = []
        if conf != "HIGH":
            issues.append(f"Company row identified with LOW confidence ({note}) - verify {agr}")

        if not rows:
            issues.append("No accused parties extracted from the petition")

        for i, r in enumerate(rows):
            if i not in assign:
                party_rows.append(_party_row(p, r, agr, None, "NOT IN COMPANY DATA", None, None,
                                             "No company borrower/co-borrower has a similar name."))
                issues.append(f"Accused {r['Party Number']} ({r['Name']}) not in company data")
                continue
            j, ns, a_s = assign[i]
            c = cparties[j]
            st = party_status(ns, a_s)
            party_rows.append(_party_row(p, r, agr, c, st, ns, a_s, _remark(st, ns, a_s)))
            if st != "MATCH":
                issues.append(f"Accused {r['Party Number']} addr {r['Address Seq']} ({r['Name']}): {st}")

        # company parties that never got paired with a petition row
        notice_names = [(norm(rows[i]["Name"]), i) for i in assign]
        assigned_slot = {i: cparties[j]["slot"] for i, (j, _n, _a) in assign.items()}
        for j, c in enumerate(cparties):
            if j in matched_j:
                continue
            same = sorted(
                ((similarity(rows[i]["Address"], c["addr"]), i) for (nn, i) in notice_names
                 if similarity(nn, c["name"]) >= MATCH_T),
                reverse=True,
            )
            if same and same[0][0] >= MATCH_T:
                oi = same[0][1]
                status, rm = "DUPLICATE IN COMPANY DATA", (
                    f"Same person and address already matched to '{rows[oi]['Name']}' under "
                    f"'{assigned_slot[oi]}' - likely a duplicate entry in company data, "
                    "not missing from the petition.")
            elif same:
                status, rm = "MISSING IN PETITION", (
                    f"Company lists a further address for '{rows[same[0][1]]['Name']}' that is not in "
                    "the petition (add it as 'ALSO AT ...' or confirm it was dropped).")
            else:
                status, rm = "MISSING IN PETITION", "Party is in company data but missing from the petition."
            party_rows.append({
                "Source File": p["source"], "SL No": p["sl_no"], "Agreement No": agr,
                "Petition Party No": "", "Address Seq": "", "Status": status, "Petition Name": "",
                "Company Name": c["name"], "Name Similarity %": None, "Petition Address": "",
                "Company Address": c["addr"], "Address Similarity %": None,
                "Company Slot": c["slot"], "Remarks": rm,
            })
            if status == "DUPLICATE IN COMPANY DATA":
                continue
            issues.append(f"{c['slot']} ({c['name']}): further address not in petition" if same
                          else f"{c['slot']} ({c['name']}) missing from petition")

        srow = {**head, "Agreement No": agr, "Matched By": by, "Match Confidence": conf}
        for field, nv, cv, res in cheque_checks(p, comp):
            srow[f"{field} - Petition"] = nv
            srow[f"{field} - Company"] = cv
            srow[f"{field} - Result"] = res
            if res != "OK":
                issues.append(f"{field}: {res}")
        srow["Cheques in Petition"] = len(p["cheques"])
        srow["Petition Entries"] = len(rows)
        srow["Company Parties"] = len(cparties)
        srow["Verdict"] = "OK" if not issues else "CHECK"
        srow["Issues"] = "; ".join(issues)
        summary_rows.append(srow)

    parties = pd.DataFrame(party_rows)
    summary = pd.DataFrame(summary_rows)
    lead = ["Source File", "SL No", "Agreement No", "Verdict", "Issues", "Matched By", "Match Confidence"]
    summary = summary[[c for c in lead if c in summary.columns] +
                      [c for c in summary.columns if c not in lead]]
    return parties, summary


# ----------------------------------------------------------------------------
# PART 5 - OUTPUT
# ----------------------------------------------------------------------------
def style_and_save(sheets, path):
    from openpyxl.styles import Alignment, Font, PatternFill

    green = PatternFill("solid", fgColor="C6EFCE")
    amber = PatternFill("solid", fgColor="FFEB9C")
    red = PatternFill("solid", fgColor="FFC7CE")
    colour = {"MATCH": green, "OK": green, "HIGH": green, "DUPLICATE IN COMPANY DATA": green,
              "PARTIAL MATCH": amber, "CHECK": amber, "LOW": amber, "NOT EXTRACTED": amber,
              "MISSING IN PETITION": amber, "MISMATCH": red, "NOT IN COMPANY DATA": red,
              "NOT FOUND": red, "NONE": red}

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(df.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(300)])
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(12, width + 2), 55)
                ws.cell(1, i).font = Font(bold=True)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top",
                                               wrap_text=isinstance(cell.value, str) and len(cell.value) > 40)
                    if isinstance(cell.value, str) and cell.value in colour:
                        cell.fill = colour[cell.value]


def save_safely(fn, path):
    try:
        fn(path)
        return path
    except PermissionError:
        alt = path.replace(".xlsx", f"_{int(time.time())}.xlsx")
        fn(alt)
        print(f"NOTE: '{path}' is open in Excel - saved as {alt}")
        return alt


# ----------------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------------
def main():
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description="Extract complaint petitions and reconcile with company data")
    ap.add_argument("--pdf-folder", default=here)
    ap.add_argument("--company", default=os.path.join(here, "company.xlsx"))
    ap.add_argument("--out", default=os.path.join(here, "petition_reconciliation_report.xlsx"))
    ap.add_argument("--extracted-out", default=os.path.join(here, "petition_extracted.xlsx"))
    args = ap.parse_args()

    pdfs = sorted(f for f in os.listdir(args.pdf_folder) if f.lower().endswith(".pdf"))
    print(f"Found {len(pdfs)} PDF file(s)\n")

    petitions = []
    for f in pdfs:
        try:
            res = extract_petitions_from_pdf(os.path.join(args.pdf_folder, f))
            if res is None:
                print(f"SKIP  [{f}]: not a complaint petition")
                continue
            for p in res:
                petitions.append(p)
                n_addr = sum(len(x["addresses"]) for x in p["parties"])
                extra = n_addr - len(p["parties"])
                msg = f"OK    [{f}] SL {p['sl_no']}: {len(p['parties'])} accused, {n_addr} address rows" \
                      + (f" ({extra} 'ALSO AT')" if extra else "") \
                      + f", {len(p['cheques'])} cheque(s)"
                print(msg)
                if not p["parties"]:
                    print(f"WARN  [{f}]: no accused found - check the 'Versus ... Accused' block")
                if not p["cheques"]:
                    print(f"WARN  [{f}]: no cheque details found")
        except Exception as e:
            print(f"ERROR [{f}]: {e}")

    if not petitions:
        sys.exit("\nNo petitions extracted.")

    extracted = flatten(petitions)
    save_safely(lambda p_: extracted.to_excel(p_, index=False), args.extracted_out)

    print(f"\nLoading company data: {args.company}")
    companies = load_company(args.company)
    print(f"  {len(companies)} agreements loaded")

    parties, summary = reconcile(petitions, companies)
    out = save_safely(
        lambda p_: style_and_save(
            {"Summary": summary, "Party Comparison": parties, "Extracted": extracted}, p_),
        args.out,
    )

    print("\nVerdicts:")
    print(summary["Verdict"].value_counts().to_string())
    print(f"\nReport: {out}")


if __name__ == "__main__":
    main()
